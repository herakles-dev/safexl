"""add_sheet on workbooks whose relationships prefix is declared per <sheet>, on the root, or under another prefix."""
from __future__ import annotations

import re
import zipfile

import openpyxl
import pytest

from safexl_mcp.structural import add_sheet

from _pkg import make_package, read_part, sheet_xml

NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def test_add_sheet_on_openpyxl_workbook(tmp_path):
    src = tmp_path / "src.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "First"
    wb.active["A1"] = "hello"
    wb.create_sheet("Second")["B2"] = 5
    wb.save(src)

    out = tmp_path / "out.xlsx"
    add_sheet(src, out, "Added")

    re_opened = openpyxl.load_workbook(out)
    assert re_opened.sheetnames == ["First", "Second", "Added"]
    assert re_opened["First"]["A1"].value == "hello"
    assert re_opened["Second"]["B2"].value == 5


def _make_with_prefix(tmp_path, prefix, on_root):
    p = make_package(tmp_path / "src.xlsx", [("One", sheet_xml('<row r="1"><c r="A1"><v>1</v></c></row>'))])
    wb = read_part(p, "xl/workbook.xml").replace("r:id", f"{prefix}:id").replace("xmlns:r=", f"xmlns:{prefix}=")
    if not on_root:
        wb = wb.replace(f' xmlns:{prefix}="{NS_R}"', "", 1)
        wb = re.sub(r"<sheet ", f'<sheet xmlns:{prefix}="{NS_R}" ', wb)
    dst = tmp_path / "pfx.xlsx"
    with zipfile.ZipFile(p) as a, zipfile.ZipFile(dst, "w") as b:
        for item in a.infolist():
            b.writestr(item, wb.encode() if item.filename == "xl/workbook.xml" else a.read(item.filename))
    return dst


# hand-built workbook.xml shapes: prefix on the root, a different prefix, and declared on each <sheet>
@pytest.mark.parametrize("prefix,on_root", [("r", True), ("rel", True), ("rel", False), ("r", False)])
def test_add_sheet_any_prefix_placement(tmp_path, prefix, on_root):
    src = _make_with_prefix(tmp_path, prefix, on_root)
    out = tmp_path / "out.xlsx"
    add_sheet(src, out, "Fresh")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["One", "Fresh"]
    assert wb["One"]["A1"].value == 1
