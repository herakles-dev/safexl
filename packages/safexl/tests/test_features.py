"""safexl_mcp.features -- feature inventory and part hashes (pure read, bounded)."""
from __future__ import annotations

import zipfile

import pytest

from safexl_mcp import features
from safexl_mcp.features import fingerprint, lost_features, untouchable_changed

from _util import fixture_path, zip_part_hashes


def test_fingerprint_counts_chart_and_drawing():
    fp = fingerprint(fixture_path("WithChart.xlsx"))
    assert fp["valid_zip"] is True and fp["error"] is None
    assert fp["features"]["charts"] >= 1
    assert fp["features"]["drawings"] >= 1


def test_fingerprint_counts_vba_pivot_and_validation():
    assert fingerprint(fixture_path("SimpleMacro.xlsm"))["features"]["vba"] == 1
    pv = fingerprint(fixture_path("ExcelPivotTableSample.xlsx"))["features"]
    assert pv["pivot_tables"] >= 1 and pv["pivot_caches"] >= 1
    assert fingerprint(fixture_path("DataValidations-49244.xlsx"))["features"]["data_validations"] >= 1


def test_fingerprint_counts_conditional_formats_including_x14():
    assert fingerprint(fixture_path("NewStyleConditionalFormattings.xlsx"))["features"]["conditional_formats"] >= 1


def test_fingerprint_counts_defined_names_and_sheets():
    f = fingerprint(fixture_path("formulas_definednames.xlsx"))["features"]
    assert f["defined_names"] >= 1 and f["sheets"] >= 1


def test_part_hashes_match_independent_sha256():
    p = fixture_path("WithChart.xlsx")
    assert fingerprint(p)["parts"] == zip_part_hashes(p)


def test_fingerprint_never_raises_on_garbage(tmp_path):
    bad = tmp_path / "x.xlsx"
    bad.write_bytes(b"not a zip at all")
    fp = fingerprint(bad)
    assert fp["valid_zip"] is False and fp["error"] and fp["parts"] == {}
    assert all(v == 0 for v in fp["features"].values())
    assert fingerprint(tmp_path / "missing.xlsx")["valid_zip"] is False


def test_part_count_cap_refused_cleanly(tmp_path, monkeypatch):
    monkeypatch.setattr(features, "MAX_PARTS", 3)
    p = tmp_path / "many.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        for i in range(5):
            z.writestr(f"p{i}.xml", "<a/>")
    fp = fingerprint(p)
    assert fp["valid_zip"] is False and "cap" in fp["error"]


def test_total_bytes_cap_refused_cleanly(tmp_path, monkeypatch):
    monkeypatch.setattr(features, "MAX_TOTAL_BYTES", 1000)
    p = tmp_path / "big.xlsx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/workbook.xml", b"A" * 5000)
    fp = fingerprint(p)
    assert fp["valid_zip"] is False and "cap" in fp["error"]


def test_caps_are_module_constants():
    assert features.MAX_PARTS > 0 and features.MAX_PART_BYTES > 0 and features.MAX_TOTAL_BYTES >= features.MAX_PART_BYTES


def test_lost_features_only_reports_decreases():
    a = {"features": {"drawings": 2, "comments": 0}}
    b = {"features": {"drawings": 1, "comments": 3}}
    assert lost_features(a, b) == [{"feature": "drawings", "before": 2, "after": 1}]


def test_untouchable_changed_flags_changed_and_removed_only():
    before = {"xl/styles.xml": "a", "docProps/core.xml": "b", "xl/worksheets/sheet1.xml": "c"}
    after = {"xl/styles.xml": "a2", "xl/worksheets/sheet1.xml": "c2"}
    assert untouchable_changed(before, after) == ["docProps/core.xml", "xl/styles.xml"]


# --- item-level losses: shifted / grown ranges are kept, moved-up / shrunk / vanished ones are lost ---

def _rng(ref, sig=("classic", "list", "")):
    return {"sig": list(sig), "ref": ref, "rects": features._rects(ref)}


@pytest.mark.parametrize("orig,new,lost", [
    ("A5:A10", "A5:A10", 0),     # untouched
    ("A5:A10", "A7:A12", 0),     # insert-rows above: shifted down
    ("A5:A10", "A5:A12", 0),     # insert-rows inside: grown
    ("A5:A10", "A3:A8", 1),      # moved up
    ("A5:A10", "A5:A8", 1),      # shrank
    ("A5:C10", "A5:B10", 1),     # narrower columns
    ("A5:A10", None, 1),         # removed
])
def test_shift_vs_removal_rule(orig, new, lost):
    got = features._unmatched_ranged([_rng(orig)], [_rng(new)] if new else [])
    assert len(got) == lost


def test_ranged_match_does_not_steal_a_later_items_counterpart():
    a, b = _rng("C36:D36", ("x",)), _rng("C60:D60", ("x",))
    assert features._unmatched_ranged([a, b], [_rng("C60:D60", ("x",)), _rng("C36:D36", ("x",))]) == []
    assert len(features._unmatched_ranged([a, b], [_rng("C60:D60", ("x",))])) == 1


def test_openpyxl_round_trip_names_item_level_losses(tmp_path):
    import openpyxl
    src = fixture_path("DataValidations-49244.xlsx")
    wb = openpyxl.load_workbook(src)
    out = tmp_path / "o.xlsx"
    wb.ws = None
    for ws in wb.worksheets:
        ws.data_validations.dataValidation = []
    wb.save(out)
    lost = lost_features(fingerprint(src), fingerprint(out))
    dv = [x for x in lost if x["feature"] == "data_validations"]
    assert dv and dv[0]["detail"] and "removed" in dv[0]["detail"][0]


def test_vba_modified_bytes_is_a_loss(tmp_path):
    src = fixture_path("SimpleMacro.xlsm")
    out = tmp_path / "m.xlsm"
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zi.infolist():
            data = zi.read(info)
            if info.filename == "xl/vbaProject.bin":
                data = data + b"\x00"
            zo.writestr(info, data)
    vba = [x for x in lost_features(fingerprint(src), fingerprint(out)) if x["feature"] == "vba"]
    assert vba and "changed" in vba[0]["detail"][0]


def test_sheet_protection_removed_is_named_with_sheet(tmp_path):
    import openpyxl
    p = tmp_path / "p.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Locked"
    wb.active.protection.sheet = True
    wb.save(p)
    q = tmp_path / "q.xlsx"
    wb2 = openpyxl.load_workbook(p)
    wb2["Locked"].protection.sheet = False
    wb2.save(q)
    prot = [x for x in lost_features(fingerprint(p), fingerprint(q)) if x["feature"] == "protection"]
    assert prot and "Locked" in prot[0]["detail"][0]


def test_defined_name_and_chart_refs_follow_a_sheet_rename():
    pairs = {"Old Name": "New"}
    assert features._norm_ref("'Old Name'!$A$1:$B$9", pairs) == features._norm_ref("New!$A$1:$B$12", {})
    assert features._norm_ref("'Old Name'!$A$1", {}) != features._norm_ref("New!$A$1", {})


def test_equivalent_boolean_spellings_are_not_a_change():
    assert features._norm_ref("TRUE()", {}) == features._norm_ref("TRUE", {})
