"""claude-plugin guard hook: driven the way Claude Code drives it (documented stdin JSON, JSON on stdout)."""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from _util import PACKAGE_DIR, fixture_path, first_sheet_name, sha256_of

PLUGIN = PACKAGE_DIR / "claude-plugin"
HOOK = PLUGIN / "hooks" / "guard_hook.py"
SAFEXL = Path(sys.executable).parent / "safexl"


def _env(tmp_path, safexl="real", **extra):
    env = {k: v for k, v in os.environ.items() if not k.startswith(("CLAUDE_", "SAFEXL_"))}
    env["CLAUDE_PLUGIN_DATA"] = str(tmp_path / "data")
    if safexl == "real":
        env["SAFEXL_BIN"] = str(SAFEXL)
    elif safexl == "missing":
        env["PATH"] = str(tmp_path / "emptybin")
    elif safexl:
        env["SAFEXL_BIN"] = str(safexl)
    env.update(extra)
    return env


def _fire(tmp_path, project, event, env, tool="Bash"):
    payload = {"session_id": "sess-1", "cwd": str(project), "hook_event_name": event,
               "transcript_path": "/dev/null", "permission_mode": "default"}
    if event == "PostToolUse":
        payload.update(tool_name=tool, tool_input={"command": "x"}, tool_response={"stdout": "", "stderr": ""})
    else:
        payload["source"] = "startup"
    mode = "session-start" if event == "SessionStart" else "post-tool"
    r = subprocess.run([sys.executable, str(HOOK), mode], input=json.dumps(payload), capture_output=True,
                       text=True, env=env, timeout=120)
    assert r.returncode == 0, r.stderr
    return json.loads(r.stdout) if r.stdout.strip() else None


@pytest.fixture
def project(tmp_path):
    p = tmp_path / "proj"
    p.mkdir()
    shutil.copy(fixture_path("WithChart.xlsx"), p / "work.xlsx")
    return p


def _context(out, event):
    spec = out["hookSpecificOutput"]
    assert spec["hookEventName"] == event
    return spec["additionalContext"]


def test_plugin_layout_declares_hooks():
    hooks = json.loads((PLUGIN / "hooks" / "hooks.json").read_text())["hooks"]
    assert set(hooks) == {"SessionStart", "PostToolUse"}
    for groups in hooks.values():
        for g in groups:
            for h in g["hooks"]:
                assert h["type"] == "command" and "${CLAUDE_PLUGIN_ROOT}/hooks/guard_hook.py" in h["args"][0]
    assert (PLUGIN / ".claude-plugin" / "plugin.json").is_file()
    assert (PLUGIN / "skills" / "safexl" / "SKILL.md").is_file()


def test_session_start_snapshots_outside_project_and_announces(tmp_path, project):
    env = _env(tmp_path)
    out = _fire(tmp_path, project, "SessionStart", env)
    assert "work.xlsx" in _context(out, "SessionStart")
    assert (tmp_path / "data" / "sessions" / "sess-1" / "state.json").is_file()
    assert sorted(p.name for p in project.iterdir()) == ["work.xlsx"]  # nothing written in the project


def test_clean_safexl_edit_is_silent(tmp_path, project):
    env = _env(tmp_path)
    _fire(tmp_path, project, "SessionStart", env)
    sheet = first_sheet_name(project / "work.xlsx")
    out = project / "step1.xlsx"
    r = subprocess.run([str(SAFEXL), "edit", str(project / "work.xlsx"), str(out), "--set", f"{sheet}!Z1=5"],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    shutil.move(out, project / "work.xlsx")
    assert _fire(tmp_path, project, "PostToolUse", env) is None


def test_openpyxl_resave_is_restored_and_model_is_told(tmp_path, project):
    env = _env(tmp_path)
    orig = sha256_of((project / "work.xlsx").read_bytes())
    _fire(tmp_path, project, "SessionStart", env)
    openpyxl.load_workbook(project / "work.xlsx").save(project / "work.xlsx")
    assert sha256_of((project / "work.xlsx").read_bytes()) != orig
    out = _fire(tmp_path, project, "PostToolUse", env, tool="Write")
    ctx = _context(out, "PostToolUse")
    assert "SAFEXL GUARD" in ctx and "work.xlsx" in ctx and "safexl" in ctx
    assert sha256_of((project / "work.xlsx").read_bytes()) == orig
    assert _fire(tmp_path, project, "PostToolUse", env) is None  # restored file is clean


def test_workbook_created_mid_session_is_tracked_and_moves_are_not_resurrected(tmp_path, project):
    env = _env(tmp_path)
    _fire(tmp_path, project, "SessionStart", env)
    new = project / "later.xlsx"
    shutil.copy(fixture_path("WithChart.xlsx"), new)
    assert _fire(tmp_path, project, "PostToolUse", env) is None  # first seen: tracked, not judged
    orig = sha256_of(new.read_bytes())
    openpyxl.load_workbook(new).save(new)
    ctx = _context(_fire(tmp_path, project, "PostToolUse", env), "PostToolUse")
    assert "later.xlsx" in ctx and sha256_of(new.read_bytes()) == orig
    new.unlink()  # deleted on purpose: never brought back
    assert _fire(tmp_path, project, "PostToolUse", env) is None
    assert not new.exists()


def test_no_workbook_means_no_output_and_no_state(tmp_path):
    empty = tmp_path / "empty"
    empty.mkdir()
    env = _env(tmp_path)
    assert _fire(tmp_path, empty, "SessionStart", env) is None
    assert _fire(tmp_path, empty, "PostToolUse", env) is None
    assert not (tmp_path / "data").exists()


def test_missing_safexl_is_nonfatal_and_says_so_once(tmp_path, project):
    (tmp_path / "emptybin").mkdir()
    env = _env(tmp_path, safexl="missing")
    first = _fire(tmp_path, project, "SessionStart", env)
    assert "safexl" in first["systemMessage"] and "PATH" in first["systemMessage"]
    assert _fire(tmp_path, project, "PostToolUse", env) is None
    assert _fire(tmp_path, project, "PostToolUse", env) is None


def test_unchanged_call_does_no_work(tmp_path, project):
    log = tmp_path / "calls.log"
    shim = tmp_path / "safexl-shim"
    shim.write_text(f'#!/bin/sh\necho "$@" >> {log}\nexec {SAFEXL} "$@"\n')
    shim.chmod(0o755)
    env = _env(tmp_path, safexl=shim)
    _fire(tmp_path, project, "SessionStart", env)
    state = tmp_path / "data" / "sessions" / "sess-1" / "state.json"
    calls_after_start = log.read_text().splitlines()
    mtime = state.stat().st_mtime_ns
    for _ in range(3):
        assert _fire(tmp_path, project, "PostToolUse", env) is None
    assert log.read_text().splitlines() == calls_after_start  # safexl never invoked again
    assert state.stat().st_mtime_ns == mtime  # state not even rewritten
    openpyxl.load_workbook(project / "work.xlsx").save(project / "work.xlsx")
    assert _fire(tmp_path, project, "PostToolUse", env) is not None
    assert len(log.read_text().splitlines()) > len(calls_after_start)  # a change does trigger the check


def test_guard_can_be_turned_off(tmp_path, project):
    assert _fire(tmp_path, project, "SessionStart", _env(tmp_path, SAFEXL_GUARD="off")) is None
    (project / ".safexl-guard-off").write_text("")
    assert _fire(tmp_path, project, "SessionStart", _env(tmp_path)) is None


def test_garbage_stdin_and_broken_safexl_never_fail(tmp_path, project):
    env = _env(tmp_path)
    r = subprocess.run([sys.executable, str(HOOK), "post-tool"], input="not json", capture_output=True,
                       text=True, env=env, cwd=project)
    assert r.returncode == 0
    bad = tmp_path / "bad-safexl"
    bad.write_text("#!/bin/sh\nexit 3\n")
    bad.chmod(0o755)
    assert _fire(tmp_path, project, "SessionStart", _env(tmp_path, safexl=bad)) is None


# ---- round 2: salvage first, restore as fallback; damaged copies of tracked workbooks -------------

def _cell(path, ref):
    sheet = first_sheet_name(path)
    out = subprocess.run([str(SAFEXL), "view", str(path), "--sheet", sheet, "--range", ref, "--json"],
                         capture_output=True, text=True)
    cells = json.loads(out.stdout)["cells"]
    return cells[0]["value"] if cells else None


def _verify_rc(before, after):
    return subprocess.run([str(SAFEXL), "verify", str(before), str(after)], capture_output=True).returncode


def _resave_with_edit(src, dst):
    wb = openpyxl.load_workbook(src)
    wb.worksheets[0]["Z1"] = 7
    wb.save(dst)


def _no_salvage_shim(tmp_path):
    shim = tmp_path / "safexl-nosalvage"
    shim.write_text(f'#!/bin/sh\nif [ "$1" = salvage ]; then echo "unsupported: test reason" >&2; exit 11; fi\n'
                    f'exec {SAFEXL} "$@"\n')
    shim.chmod(0o755)
    return shim


@pytest.fixture
def drawing_project(tmp_path):
    p = tmp_path / "proj"
    p.mkdir()
    shutil.copy(fixture_path("WithDrawing.xlsx"), p / "work.xlsx")
    return p


def test_salvage_rebuilds_from_original_and_keeps_the_edit(tmp_path, drawing_project):
    env = _env(tmp_path)
    keep = tmp_path / "orig.xlsx"
    shutil.copy(drawing_project / "work.xlsx", keep)
    _fire(tmp_path, drawing_project, "SessionStart", env)
    _resave_with_edit(drawing_project / "work.xlsx", drawing_project / "work.xlsx")
    assert _verify_rc(keep, drawing_project / "work.xlsx") == 11
    ctx = _context(_fire(tmp_path, drawing_project, "PostToolUse", env), "PostToolUse")
    assert "rebuilt from the original" in ctx and "Z1" in ctx and "safexl" in ctx
    assert _verify_rc(keep, drawing_project / "work.xlsx") == 0
    assert _cell(drawing_project / "work.xlsx", "Z1") == 7
    assert _fire(tmp_path, drawing_project, "PostToolUse", env) is None
    assert sorted(p.name for p in drawing_project.iterdir()) == ["work.xlsx"]


def test_salvage_unsupported_falls_back_to_restore_with_reason(tmp_path, drawing_project):
    env = _env(tmp_path, safexl=_no_salvage_shim(tmp_path))
    orig = sha256_of((drawing_project / "work.xlsx").read_bytes())
    _fire(tmp_path, drawing_project, "SessionStart", env)
    _resave_with_edit(drawing_project / "work.xlsx", drawing_project / "work.xlsx")
    ctx = _context(_fire(tmp_path, drawing_project, "PostToolUse", env), "PostToolUse")
    assert "unsupported: test reason" in ctx and "Redo the edit" in ctx
    assert sha256_of((drawing_project / "work.xlsx").read_bytes()) == orig


def test_restore_mode_keeps_round_one_behaviour(tmp_path, drawing_project):
    env = _env(tmp_path, SAFEXL_GUARD_MODE="restore")
    orig = sha256_of((drawing_project / "work.xlsx").read_bytes())
    _fire(tmp_path, drawing_project, "SessionStart", env)
    _resave_with_edit(drawing_project / "work.xlsx", drawing_project / "work.xlsx")
    ctx = _context(_fire(tmp_path, drawing_project, "PostToolUse", env), "PostToolUse")
    assert "rebuilt from the original" not in ctx and "Redo the edit" in ctx
    assert sha256_of((drawing_project / "work.xlsx").read_bytes()) == orig


def test_new_derivative_file_is_repaired_and_original_untouched(tmp_path, drawing_project):
    env = _env(tmp_path)
    orig_sha = sha256_of((drawing_project / "work.xlsx").read_bytes())
    _fire(tmp_path, drawing_project, "SessionStart", env)
    _resave_with_edit(drawing_project / "work.xlsx", drawing_project / "report.xlsx")
    ctx = _context(_fire(tmp_path, drawing_project, "PostToolUse", env), "PostToolUse")
    assert "report.xlsx" in ctx and "rebuilt from the original" in ctx and "original workbook was not touched" in ctx
    assert _verify_rc(drawing_project / "work.xlsx", drawing_project / "report.xlsx") == 0
    assert _cell(drawing_project / "report.xlsx", "Z1") == 7
    assert sha256_of((drawing_project / "work.xlsx").read_bytes()) == orig_sha
    assert _fire(tmp_path, drawing_project, "PostToolUse", env) is None


def test_derivative_that_cannot_be_salvaged_is_left_with_instructions(tmp_path, drawing_project):
    env = _env(tmp_path, safexl=_no_salvage_shim(tmp_path))
    _fire(tmp_path, drawing_project, "SessionStart", env)
    _resave_with_edit(drawing_project / "work.xlsx", drawing_project / "report.xlsx")
    damaged = sha256_of((drawing_project / "report.xlsx").read_bytes())
    ctx = _context(_fire(tmp_path, drawing_project, "PostToolUse", env), "PostToolUse")
    assert "cp <original> report.xlsx" in ctx and "unsupported: test reason" in ctx
    assert sha256_of((drawing_project / "report.xlsx").read_bytes()) == damaged
    assert _fire(tmp_path, drawing_project, "PostToolUse", env) is None  # said once for this exact file


def test_new_unrelated_workbook_is_tracked_silently(tmp_path, drawing_project):
    env = _env(tmp_path)
    _fire(tmp_path, drawing_project, "SessionStart", env)
    other = drawing_project / "other.xlsx"
    shutil.copy(fixture_path("ExcelPivotTableSample.xlsx"), other)
    assert _fire(tmp_path, drawing_project, "PostToolUse", env) is None
    state = json.loads((tmp_path / "data" / "sessions" / "sess-1" / "state.json").read_text())
    assert str(other) in state["tracked"]


def test_lossless_resave_is_rebuilt_with_edit_not_restored(tmp_path, project):
    env = _env(tmp_path)
    keep = tmp_path / "orig.xlsx"
    shutil.copy(project / "work.xlsx", keep)
    _fire(tmp_path, project, "SessionStart", env)
    _resave_with_edit(project / "work.xlsx", project / "work.xlsx")  # WithChart: rewritten, nothing lost
    assert _verify_rc(keep, project / "work.xlsx") == 10
    ctx = _context(_fire(tmp_path, project, "PostToolUse", env), "PostToolUse")
    assert "re-saved by another library" in ctx and "byte-identical" in ctx and "Z1" in ctx
    assert "lost" not in ctx and "restored" not in ctx.lower()
    assert _verify_rc(keep, project / "work.xlsx") == 0
    assert _cell(project / "work.xlsx", "Z1") == 7


def test_restore_mode_still_restores_a_lossless_resave(tmp_path, project):
    env = _env(tmp_path, SAFEXL_GUARD_MODE="restore")
    orig = sha256_of((project / "work.xlsx").read_bytes())
    _fire(tmp_path, project, "SessionStart", env)
    _resave_with_edit(project / "work.xlsx", project / "work.xlsx")
    ctx = _context(_fire(tmp_path, project, "PostToolUse", env), "PostToolUse")
    assert "restored" in ctx and "Redo the edit" in ctx
    assert sha256_of((project / "work.xlsx").read_bytes()) == orig
