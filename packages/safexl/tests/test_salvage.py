"""Tests for safexl_mcp.salvage: re-apply a damaged copy's intended edits to the original."""
from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.chart import BarChart, Reference  # noqa: E402

import safexl_mcp.salvage as salvage_mod  # noqa: E402
from safexl_mcp.engine import edit  # noqa: E402
from safexl_mcp.salvage import salvage  # noqa: E402


def _hashes(p: Path) -> dict[str, str]:
    with zipfile.ZipFile(p) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}


def _make_original(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["name", "qty", "total"])
    for i in range(2, 7):
        ws.append([f"item{i}", i * 10, None])
        ws.cell(row=i, column=3).value = f"=B{i}*2"
    other = wb.create_sheet("Other")
    other["A1"] = "keep"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    ws.add_chart(chart, "F2")
    wb.save(path)
    _hoist_r_namespace(path)
    return path


def _hoist_r_namespace(path: Path) -> None:
    """openpyxl declares xmlns:r on every <sheet>; Excel declares it on the root,
    which is what structural.add_sheet's appended ``r:id`` relies on."""
    ns = ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "xl/workbook.xml":
                text = data.decode().replace(ns, "")
                data = text.replace("<workbook ", "<workbook" + ns + " ", 1).encode()
            zo.writestr(item, data)
    tmp.replace(path)


def _resave(src: Path, dst: Path, mutate=None) -> Path:
    wb = openpyxl.load_workbook(src)
    if mutate:
        mutate(wb)
    wb.save(dst)
    return dst


def test_two_edits_salvaged_and_untouched_parts_byte_identical(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")

    def mutate(wb):
        wb["Data"]["B3"] = 999
        wb["Other"]["A1"] = "changed"

    damaged = _resave(orig, tmp_path / "dam.xlsx", mutate)
    assert _hashes(damaged) != _hashes(orig)  # the re-save rewrote parts

    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out, force=True)
    assert res["status"] == "salvaged"
    assert {(a["sheet"], a["ref"]) for a in res["applied"]} == {("Data", "B3"), ("Other", "A1")}
    assert res["skipped"] == []

    ho, hs = _hashes(orig), _hashes(out)
    assert any(n.startswith("xl/charts/") for n in hs)
    changed = {n for n in ho if ho[n] != hs.get(n)}
    # only the two edited sheets and the engine's own workbook-level bookkeeping change
    assert changed <= {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml", "xl/workbook.xml", "xl/calcChain.xml", "[Content_Types].xml", "xl/_rels/workbook.xml.rels"}
    assert all(ho[n] == hs[n] for n in ho if n.startswith(("xl/charts/", "xl/drawings/", "xl/styles", "xl/theme")))
    assert not any(p.startswith("xl/charts/") for p in res["parts"]["changed"])

    wb = openpyxl.load_workbook(out)
    assert wb["Data"]["B3"].value == 999
    assert wb["Other"]["A1"].value == "changed"
    assert wb["Data"]["C3"].value == "=B3*2"


def test_pure_resave_is_nothing_to_do(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")
    damaged = _resave(orig, tmp_path / "dam.xlsx")
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out)
    assert res["status"] == "nothing_to_do"
    assert res["applied"] == [] and res["skipped"] == []
    assert not out.exists()


def test_formula_edit_applied(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")
    damaged = _resave(orig, tmp_path / "dam.xlsx", lambda wb: wb["Data"].__setitem__("C4", "=B4*3"))
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out, force=True)
    assert res["status"] == "salvaged"
    assert res["applied"] == [{"sheet": "Data", "ref": "C4", "kind": "formula"}]
    assert openpyxl.load_workbook(out)["Data"]["C4"].value == "=B4*3"


def test_removed_sheet_detected_and_not_misapplied(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")

    def mutate(wb):
        wb["Data"]["B3"] = 5
        del wb["Other"]

    damaged = _resave(orig, tmp_path / "dam.xlsx", mutate)
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out, force=True)
    assert res["status"] == "unsupported"
    assert {"kind": "sheet_removed", "sheet": "Other", "handled": False} in res["structural"]["detected"]
    assert not out.exists()


def test_column_shift_detected_unsupported(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")
    damaged = _resave(orig, tmp_path / "dam.xlsx", lambda wb: wb["Data"].insert_cols(2))
    res = salvage(orig, damaged, tmp_path / "out.xlsx", force=True)
    assert res["status"] == "unsupported"
    assert any(d["kind"] == "columns_shifted" and not d["handled"] for d in res["structural"]["detected"])
    assert not (tmp_path / "out.xlsx").exists()


def test_uniform_row_insert_handled(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")

    def mutate(wb):
        ws = wb["Data"]
        ws.insert_rows(4)
        ws["A4"] = "new"
        ws["B4"] = 77

    damaged = _resave(orig, tmp_path / "dam.xlsx", mutate)
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out, force=True)
    assert {"kind": "rows_inserted", "sheet": "Data", "first_row": 4, "count": 1, "handled": True} in res["structural"]["detected"]
    wb = openpyxl.load_workbook(out)
    ws = wb["Data"]
    assert (ws["A4"].value, ws["B4"].value) == ("new", 77)
    assert ws["A5"].value == "item4"
    # openpyxl left the old formulas stale; the engine's own adjusted ones must win
    assert ws["C5"].value == "=B5*2"
    assert any(n.startswith("xl/charts/") for n in _hashes(out))


def test_rename_and_added_sheet_handled(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")

    def mutate(wb):
        wb["Other"].title = "Renamed"
        ws = wb.create_sheet("Check")
        ws["A1"] = "hello"
        ws["B1"] = "=Data!B2"

    damaged = _resave(orig, tmp_path / "dam.xlsx", mutate)
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out, force=True)
    kinds = {d["kind"] for d in res["structural"]["detected"]}
    assert {"sheet_renamed", "sheet_added"} <= kinds
    assert res["status"] == "salvaged"
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Data", "Renamed", "Check"]
    assert wb["Check"]["A1"].value == "hello" and wb["Check"]["B1"].value == "=Data!B2"
    assert any(n.startswith("xl/charts/") for n in _hashes(out))


def _strip_chart(path: Path) -> None:
    """Make the re-save lose the chart, the way a lossy writer would."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            if item.filename.startswith(("xl/charts/", "xl/drawings/")):
                continue
            data = zi.read(item.filename)
            if item.filename.endswith(".rels") and b"drawing" in data:
                data = b"".join(l for l in data.replace(b"/>", b"/>\n").splitlines(True) if b"drawing" not in l)
            if item.filename == "[Content_Types].xml":
                data = b"".join(l for l in data.replace(b"/>", b"/>\n").splitlines(True) if b"chart" not in l and b"drawing" not in l)
            if item.filename.startswith("xl/worksheets/sheet") and b"<drawing" in data:
                import re
                data = re.sub(rb"<drawing[^>]*/>", b"", data)
            zo.writestr(item, data)
    tmp.replace(path)


def test_gate_rewritten_without_loss_is_nothing_to_do_unless_forced(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")
    damaged = _resave(orig, tmp_path / "dam.xlsx", lambda wb: wb["Data"].__setitem__("B3", 5))
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out)
    assert res["status"] == "nothing_to_do" and res["gate"]["verdict"] in ("rewritten", "intact")
    assert not out.exists()
    assert salvage(orig, damaged, out, force=True)["status"] == "salvaged"


def test_lossy_resave_proceeds_by_default_and_self_check_is_intact(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")
    damaged = _resave(orig, tmp_path / "dam.xlsx", lambda wb: wb["Data"].__setitem__("B3", 5))
    _strip_chart(damaged)
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out)
    assert res["gate"]["verdict"] in ("damaged", "invalid")
    assert res["status"] == "salvaged" and res["verify"]["verdict"] == "intact"
    assert out.exists()


def test_safexl_written_file_is_nothing_to_do(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")
    written = tmp_path / "written.xlsx"
    edit(orig, written, [{"sheet": "Data", "cell": "B3", "value": 42}])
    out = tmp_path / "out.xlsx"
    res = salvage(orig, written, out)
    assert res["status"] == "nothing_to_do" and res["gate"]["verdict"] == "intact"
    assert not out.exists()


def test_table_growth_carried_over(tmp_path):
    from openpyxl.worksheet.table import Table

    orig = tmp_path / "orig.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T"
    ws.append(["a", "b"])
    ws.append([1, 2])
    ws.append([3, 4])
    ws.add_table(Table(displayName="Tbl", ref="A1:B3"))
    wb.save(orig)

    def mutate(wb):
        ws = wb["T"]
        ws["C1"], ws["C2"], ws["C3"] = "c", 5, 6
        ws.tables["Tbl"].ref = "A1:C3"

    damaged = _resave(orig, tmp_path / "dam.xlsx", mutate)
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out, force=True)
    assert res["tables"] == [{"table": "Tbl", "sheet": "T", "old_ref": "A1:B3", "ref": "A1:C3"}]
    assert res["status"] == "salvaged"
    wb = openpyxl.load_workbook(out)
    assert wb["T"].tables["Tbl"].ref == "A1:C3"
    assert [c.name for c in wb["T"].tables["Tbl"].tableColumns] == ["a", "b", "c"]


def test_table_change_not_explained_is_reported_not_applied(tmp_path):
    from openpyxl.worksheet.table import Table

    orig = tmp_path / "orig.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T"
    ws.append(["a", "b"])
    ws.append([1, 2])
    ws.add_table(Table(displayName="Tbl", ref="A1:B2"))
    ws["A5"] = "x"
    wb.save(orig)

    def mutate(wb):
        wb["T"].tables["Tbl"].ref = "A1:B3"  # grew, but no cell edit in the new row
        wb["T"]["A5"] = "y"

    damaged = _resave(orig, tmp_path / "dam.xlsx", mutate)
    res = salvage(orig, damaged, tmp_path / "out.xlsx", force=True)
    assert res["tables"] == []
    assert any("no cell edit explains" in s["reason"] for s in res["skipped"])
    assert res["status"] == "partial"


def test_column_scoped_shift_applied_with_insert_cells(tmp_path):
    orig = tmp_path / "orig.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "L"
    for i in range(1, 7):
        ws.cell(row=i, column=1, value=f"a{i}")
        ws.cell(row=i, column=3, value=f"c{i}")
    wb.save(orig)

    def mutate(wb):
        ws = wb["L"]
        ws.move_range("A3:A6", rows=1)  # only column A shifts
        ws["A3"] = "NEW"

    damaged = _resave(orig, tmp_path / "dam.xlsx", mutate)
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out, force=True)
    assert {"kind": "cells_inserted_down", "sheet": "L", "range": "A3:A3", "handled": True} in res["structural"]["detected"]
    ws = openpyxl.load_workbook(out)["L"]
    assert [ws.cell(row=r, column=1).value for r in range(1, 8)] == ["a1", "a2", "NEW", "a3", "a4", "a5", "a6"]
    assert [ws.cell(row=r, column=3).value for r in range(1, 7)] == [f"c{i}" for i in range(1, 7)]


def test_engine_refusal_is_unsupported_not_an_exception(tmp_path):
    orig = _make_original(tmp_path / "orig.xlsx")
    # make the original structurally invalid: a row whose cells are out of order
    bad = tmp_path / "bad.xlsx"
    with zipfile.ZipFile(orig) as zi, zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "xl/worksheets/sheet2.xml":
                data = data.replace(b'<c r="A1"', b'<c r="B1"><v>1</v></c><c r="A1"', 1)
            zo.writestr(item, data)
    damaged = _resave(orig, tmp_path / "dam.xlsx", lambda wb: wb["Other"].__setitem__("A1", "z"))
    out = tmp_path / "out.xlsx"
    res = salvage(bad, damaged, out, force=True)
    assert res["status"] == "unsupported"
    assert any("refused" in s["reason"] for s in res["skipped"])
    assert not out.exists()


def test_failed_self_check_deletes_output(tmp_path, monkeypatch):
    orig = _make_original(tmp_path / "orig.xlsx")
    damaged = _resave(orig, tmp_path / "dam.xlsx", lambda wb: wb["Data"].__setitem__("B3", 5))
    real = salvage_mod.verify
    calls = []

    def fake(a, b):
        calls.append(1)
        r = real(a, b)
        if len(calls) == 2:
            r = {**r, "verdict": "damaged", "reasons": ["forced"], "lost_features": []}
        return r

    monkeypatch.setattr(salvage_mod, "verify", fake)
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out, force=True)
    assert res["status"] == "unsupported" and "self-check failed" in res["skipped"][-1]["reason"]
    assert not out.exists()


def test_two_column_scoped_blocks_applied(tmp_path):
    orig = tmp_path / "orig.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "L"
    for i in range(1, 7):
        ws.cell(row=i, column=1, value=f"a{i}")
        ws.cell(row=i, column=3, value=f"c{i}")
        ws.cell(row=i, column=5, value=f"e{i}")  # never moves
    wb.save(orig)

    def mutate(wb):
        ws = wb["L"]
        ws.move_range("A3:A6", rows=1)
        ws["A3"] = "NA"
        ws.move_range("C4:C6", rows=1)  # a different insert row in column C
        ws["C4"] = "NC"

    damaged = _resave(orig, tmp_path / "dam.xlsx", mutate)
    out = tmp_path / "out.xlsx"
    res = salvage(orig, damaged, out, force=True)
    assert {d["range"] for d in res["structural"]["detected"] if d["kind"] == "cells_inserted_down"} == {"A3:A3", "C4:C4"}
    ws = openpyxl.load_workbook(out)["L"]
    assert [ws.cell(row=r, column=1).value for r in range(1, 8)] == ["a1", "a2", "NA", "a3", "a4", "a5", "a6"]
    assert [ws.cell(row=r, column=3).value for r in range(1, 8)] == ["c1", "c2", "c3", "NC", "c4", "c5", "c6"]
    assert [ws.cell(row=r, column=5).value for r in range(1, 7)] == [f"e{i}" for i in range(1, 7)]
